
from pathlib import Path
import pymupdf as fitz
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_pdf(file_path):
    """Extract text from a PDF, preserving page numbers."""

    extracted_pages = []

    with fitz.open(file_path) as pdf:

        for page_number, page in enumerate(pdf, start=1):

            text = page.get_text("text")

            if text.strip():
                extracted_pages.append({
                    "page_number": page_number,
                    "text": text
                })

    return extracted_pages


def extract_docx(file_path):
    """Extract text from paragraphs and tables in a DOCX document."""

    document = Document(file_path)
    extracted_content = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            extracted_content.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            extracted_content.append(" | ".join(cells))

    return [{
        "page_number": None,
        "text": "\n".join(extracted_content)
    }]


def extract_txt(file_path):
    """Read text from a TXT document."""

    text = Path(file_path).read_text(encoding="utf-8-sig")

    return [{
        "page_number": None,
        "text": text
    }]


def load_document(file_path):
    """
    Detect document type, extract text, and return structured data.
    """

    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file format: {extension}"
        )

    if not file_path.is_file():
        raise FileNotFoundError(
            f"Document not found: {file_path}"
        )

    if extension == ".pdf":
        pages = extract_pdf(file_path)

    elif extension == ".docx":
        pages = extract_docx(file_path)

    else:
        pages = extract_txt(file_path)

    return {
        "filename": file_path.name,
        "file_type": extension,
        "source_path": str(file_path),
        "pages": pages
    }


def load_documents(folder_path):
    """Load all supported documents from a folder."""

    folder = Path(folder_path)

    if not folder.is_dir():
        raise NotADirectoryError(
            f"Document folder not found: {folder}"
        )

    documents = []

    for file_path in sorted(folder.iterdir()):

        if not file_path.is_file():
            continue

        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            print(f"[SKIPPED] Unsupported format: {file_path.name}")
            continue

        try:
            document = load_document(file_path)

            documents.append(document)

            print(f"[SUCCESS] Loaded: {file_path.name}")

        except Exception as error:
            print(f"[ERROR] {file_path.name}: {error}")

    return documents
